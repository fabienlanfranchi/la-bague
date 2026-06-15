"""Génère un fichier Excel de sauvegarde complète de toutes les données du club."""
import os
import pandas as pd
from pymongo import MongoClient
from dotenv import load_dotenv
from datetime import datetime
from pathlib import Path

load_dotenv(Path(__file__).parent / '.env')
client = MongoClient(os.environ['MONGO_URL'])
db = client[os.environ['DB_NAME']]

members = {m['id']: m for m in db.members.find({})}
print(f"Membres: {len(members)}")

# 1. Saisons config
saisons = list(db.saisons_config.find({}, {'_id': 0}).sort('saison', 1))
df_saisons = pd.DataFrame([{
    'Saison': s.get('saison'),
    'Nb Apéros': s.get('nb_aperos', 0),
    'Nb Repas': s.get('nb_repas', 0),
    'Nb Anniversaires': s.get('nb_anniversaires', 0),
    'Présences Apéros (total)': s.get('presences_membres_aperos', 0),
    'Présences Repas (total)': s.get('presences_membres_repas', 0),
    'Présences Anniversaires (total)': s.get('presences_membres_anniversaires', 0),
    'Nb membres': s.get('nb_membres_manuel', 0),
    'Verrouillée (manuel)': 'OUI' if s.get('is_manuel') else 'NON',
} for s in saisons])

# 2. Présences par membre par saison
presences = list(db.presences_membres.find({}, {'_id': 0}))
rows = []
for p in presences:
    m = members.get(p.get('membre_id'), {})
    rows.append({
        'Saison': p.get('saison'),
        'N° Membre': m.get('numero_membre', ''),
        'Nom Complet': m.get('nom_complet', '(inconnu)'),
        'Fonction': m.get('fonction', ''),
        'Présences Apéros': p.get('presences_aperos', 0),
        'Présences Repas': p.get('presences_repas', 0),
        'Présences Anniversaires': p.get('presences_anniversaires', 0),
        'Total Présences': p.get('presences_aperos', 0) + p.get('presences_repas', 0) + p.get('presences_anniversaires', 0),
    })
df_presences = pd.DataFrame(rows).sort_values(['Saison', 'N° Membre'])

# 3. Événements
evenements = list(db.evenements.find({}, {'_id': 0}).sort([('saison', 1), ('date', 1)]))
df_evt = pd.DataFrame([{
    'ID Événement': e.get('id'),
    'Saison': e.get('saison'),
    'Date': e.get('date', '')[:10] if e.get('date') else '',
    'Objet': e.get('objet', ''),
    'Lieu': e.get('lieu', ''),
    'Type': e.get('type_sondage', ''),
    'Statut': e.get('statut', ''),
    'Total Présents (membres)': e.get('total_presents', 0),
} for e in evenements])

# 4. Réponses manuelles
rm = list(db.reponses_manuelles.find({}, {'_id': 0}))
evt_map = {e['id']: e for e in evenements}
rows_rm = []
for r in rm:
    e = evt_map.get(r.get('evenement_id'), {})
    nom = r.get('nom', '')
    if r.get('membre_id'):
        mb = members.get(r['membre_id'], {})
        nom = mb.get('nom_complet', nom)
    rows_rm.append({
        'Saison': e.get('saison', ''),
        'Date Événement': e.get('date', '')[:10] if e.get('date') else '',
        'Objet': e.get('objet', ''),
        'Type Personne': r.get('type', ''),
        'Nom': nom,
        'Présent': 'OUI' if r.get('present') else 'NON',
        'Entrée': r.get('choix_entree', ''),
        'Plat': r.get('choix_plat', ''),
        'Dessert': r.get('choix_dessert', ''),
    })
df_rm = pd.DataFrame(rows_rm).sort_values(['Saison', 'Date Événement'])

# 5. Liste des membres
df_membres = pd.DataFrame([{
    'N°': m.get('numero_membre'),
    'Nom Complet': m.get('nom_complet'),
    'Fonction': m.get('fonction', ''),
    'Email': m.get('email', ''),
    'Téléphone': m.get('telephone', ''),
    'Saison Entrée': m.get('saison_entree', ''),
    'Année Entrée': m.get('annee_entree', ''),
    'Compte Validé': 'OUI' if m.get('is_validated') else 'NON',
} for m in members.values()]).sort_values('N°')

os.makedirs('/app/exports', exist_ok=True)
ts = datetime.now().strftime('%Y%m%d_%H%M%S')
filepath = f'/app/exports/SAUVEGARDE_LaBagueImperiale_{ts}.xlsx'

with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
    df_saisons.to_excel(writer, sheet_name='Récap Saisons', index=False)
    df_presences.to_excel(writer, sheet_name='Présences par Membre', index=False)
    df_evt.to_excel(writer, sheet_name='Événements', index=False)
    df_rm.to_excel(writer, sheet_name='Présences par Événement', index=False)
    df_membres.to_excel(writer, sheet_name='Membres', index=False)

from openpyxl import load_workbook
wb = load_workbook(filepath)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 50)
wb.save(filepath)

print(f"Saisons: {len(df_saisons)}, Présences: {len(df_presences)}, Événements: {len(df_evt)}, Réponses manuelles: {len(df_rm)}, Membres: {len(df_membres)}")
size_kb = os.path.getsize(filepath) / 1024
print(f"\nFichier généré : {filepath}")
print(f"Taille : {size_kb:.1f} KB")
