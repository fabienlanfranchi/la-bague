"""
Restauration Saison 1 DEPUIS LE FICHIER xlsx OFFICIEL.
Source: /tmp/S1.xlsx (Saison_1_Bague_Imperiale_Presences.xlsx)
- 13 repas (PAS d'anniversaire !)
- 0 apéros
- "Chemin des Vignobles" reclassé en REPAS
"""
import asyncio, os, uuid
import openpyxl
from datetime import datetime, timezone
from motor.motor_asyncio import AsyncIOMotorClient
from dotenv import load_dotenv
load_dotenv("/app/backend/.env")

# Charger le xlsx
wb = openpyxl.load_workbook("/tmp/S1.xlsx", data_only=True)

# 1. SYNTHÈSE
synth = wb["Synthèse"]
NB_EVENTS = synth.cell(row=4, column=2).value      # 13
NB_APEROS = synth.cell(row=9, column=2).value      # 0
NB_REPAS = synth.cell(row=10, column=2).value      # 13
NB_ANNIV = synth.cell(row=11, column=2).value      # 0
NB_MEMBRES_INSCRITS = synth.cell(row=7, column=2).value  # 34
PRES_TOTAL = synth.cell(row=5, column=2).value     # 233
print(f"Synthèse S1: events={NB_EVENTS} (ap={NB_APEROS}, rp={NB_REPAS}, an={NB_ANNIV}), inscrits={NB_MEMBRES_INSCRITS}, prés_total={PRES_TOTAL}")

# 2. EVENTS (ordre + catégories + dates)
ev_sheet = wb["Présences par événement"]
EVENTS = []
for row in range(2, 15):
    nom = ev_sheet.cell(row=row, column=1).value
    date = ev_sheet.cell(row=row, column=2).value
    cat = ev_sheet.cell(row=row, column=3).value
    presents = ev_sheet.cell(row=row, column=4).value
    if not nom:
        break
    EVENTS.append({"nom": nom, "date": date, "categorie": cat, "presents": presents})
print(f"\n{len(EVENTS)} événements:")
for e in EVENTS:
    print(f"  - {e['nom']} ({e['date']}) | {e['categorie']} | {e['presents']} prés")

# 3. MATRICE
matrix_sheet = wb["Matrice"]
# Membres et leurs présences par event
MATRIX_BY_NAME = {}
for row in range(2, 36):
    nom_membre = matrix_sheet.cell(row=row, column=1).value
    if not nom_membre:
        continue
    presences = [matrix_sheet.cell(row=row, column=c).value for c in range(2, 15)]
    presences = [int(p or 0) for p in presences]
    MATRIX_BY_NAME[nom_membre] = presences

# Mapping nom xlsx -> numéro membre actif en base
ACTIVE_MAP = {
    "Fabien": 1, "Jacques": 3, "Nini": 4, "Ange Phi": 5, "Jeff": 6,
    "Mathias": 8, "Jean Jacques": 9, "Ludo Leca": 18, "Pascal": 23, "Paul Rossion": 24,
}

# Type de chaque colonne d'event
def cat_to_type(cat):
    c = str(cat).lower().strip()
    if c == "apéro" or c == "apéros" or c == "apero": return "apero"
    if c == "repas": return "repas"
    if c == "anniversaire" or c == "anniversaires": return "anniversaire"
    return "repas"

async def main():
    client = AsyncIOMotorClient(os.environ["MONGO_URL"])
    db = client[os.environ["DB_NAME"]]
    now_iso = datetime.now(timezone.utc).isoformat()
    
    # BACKUP
    snapshot_pres = await db.presences_membres.find({"saison": 1}, {"_id": 0}).to_list(500)
    snapshot_cfg = await db.saisons_config.find_one({"saison": 1}, {"_id": 0})
    snapshot_evts = await db.evenements.find({"saison": 1}, {"_id": 0}).to_list(100)
    backup_id = str(uuid.uuid4())
    await db.presences_membres_backup.insert_one({
        "id": backup_id, "saison": 1, "backup_date": now_iso,
        "reason": "manual_restore_S1_from_xlsx_official",
        "presences_snapshot": snapshot_pres, "config_snapshot": snapshot_cfg, "events_snapshot": snapshot_evts
    })
    print(f"\n✅ BACKUP id={backup_id}")
    
    # 1) saisons_config.S1
    await db.saisons_config.update_one(
        {"saison": 1},
        {"$set": {
            "nb_aperos": NB_APEROS, "nb_repas": NB_REPAS, "nb_anniversaires": NB_ANNIV,
            "is_manuel": True,
            "presences_membres_aperos": 0,
            "presences_membres_repas": PRES_TOTAL,
            "presences_membres_anniversaires": 0,
            "nb_membres_manuel": NB_MEMBRES_INSCRITS,
            "updated_at": now_iso,
            "restored_from_xlsx_15_06_2026": True
        }}
    )
    print(f"✅ saisons_config S1 -> ap={NB_APEROS}, rp={NB_REPAS}, an={NB_ANNIV}, nb_membres={NB_MEMBRES_INSCRITS}, prés_repas_total={PRES_TOTAL}")
    
    # 2) Identifier et MAJ les events S1
    base_evts = await db.evenements.find({"saison": 1}, {"_id": 0}).sort("date", 1).to_list(50)
    # Mapping par lieu (fuzzy)
    def normalize(s):
        return (s or "").lower().replace("'", "").replace("é", "e").replace("è", "e").replace("ê", "e").strip()
    
    ordered_base = []
    for ev_x in EVENTS:
        target_nom = normalize(ev_x["nom"])
        matched = None
        for b in base_evts:
            if normalize(b["lieu"]) == target_nom:
                # Préférer celui pas encore matché
                if b not in [o["base"] for o in ordered_base]:
                    matched = b
                    break
        ordered_base.append({"x": ev_x, "base": matched})
    
    # Update each event type
    for o in ordered_base:
        if not o["base"]:
            print(f"⚠️ Event introuvable en base: {o['x']['nom']}")
            continue
        new_type = cat_to_type(o["x"]["categorie"])
        await db.evenements.update_one(
            {"id": o["base"]["id"]},
            {"$set": {
                "type_sondage": new_type,
                "total_presents": o["x"]["presents"],
                "statut": "terminé"
            }}
        )
        print(f"  Event '{o['x']['nom']}' -> type={new_type}, total={o['x']['presents']} (was: {o['base'].get('type_sondage')})")
    
    # 3) presences_membres pour les membres actifs
    members = await db.members.find({"numero_membre": {"$in": list(ACTIVE_MAP.values())}}, {"_id": 0}).to_list(50)
    md_by_num = {m["numero_membre"]: m for m in members}
    
    for xlsx_name, num in ACTIVE_MAP.items():
        m = md_by_num.get(num)
        if not m:
            print(f"⚠️ Membre #{num} ({xlsx_name}) introuvable en base")
            continue
        row = MATRIX_BY_NAME.get(xlsx_name)
        if not row:
            print(f"⚠️ {xlsx_name} pas dans matrice")
            continue
        # Calculer par catégorie
        ap = rp = an = 0
        for i, ev_x in enumerate(EVENTS):
            if row[i] != 1:
                continue
            t = cat_to_type(ev_x["categorie"])
            if t == "apero": ap += 1
            elif t == "repas": rp += 1
            elif t == "anniversaire": an += 1
        await db.presences_membres.update_one(
            {"membre_id": m["id"], "saison": 1},
            {"$set": {
                "presences_aperos": ap, "presences_repas": rp, "presences_anniversaires": an,
                "updated_at": now_iso, "restored_from_xlsx_15_06_2026": True
            }},
            upsert=True
        )
        print(f"  #{num:>2} {m['nom_complet'][:30]:<32} -> ap={ap} rp={rp} an={an}")
    
    # 4) Recréer les reponses_manuelles selon la matrice
    s1_evt_ids = [o["base"]["id"] for o in ordered_base if o["base"]]
    member_ids = [md_by_num[v]["id"] for v in ACTIVE_MAP.values() if v in md_by_num]
    r = await db.reponses_manuelles.delete_many({
        "evenement_id": {"$in": s1_evt_ids},
        "membre_id": {"$in": member_ids}
    })
    print(f"\n🗑️ {r.deleted_count} reponses_manuelles anciennes supprimées")
    
    total_created = 0
    for i, o in enumerate(ordered_base):
        if not o["base"]: continue
        for xlsx_name, num in ACTIVE_MAP.items():
            row = MATRIX_BY_NAME.get(xlsx_name, [])
            if i >= len(row) or row[i] != 1:
                continue
            m = md_by_num.get(num)
            if not m: continue
            await db.reponses_manuelles.insert_one({
                "id": str(uuid.uuid4()), "evenement_id": o["base"]["id"], "membre_id": m["id"],
                "nom": m["nom_complet"], "type": "membre_manuel", "present": True,
                "choix_entree": None, "choix_plat": None, "choix_dessert": None,
                "created_at": now_iso, "restored_from_xlsx_15_06_2026": True
            })
            total_created += 1
    print(f"✅ {total_created} reponses_manuelles créées")
    
    client.close()

asyncio.run(main())
